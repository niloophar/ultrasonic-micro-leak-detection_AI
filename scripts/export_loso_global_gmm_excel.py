#!/usr/bin/env python
"""
Export the LOSO Global-GMM results (loso_results.json + majority_vote_results.json)
to a single formatted Excel workbook for the thesis writeup.

Reads:
  data/results/interpolation_global_gmm/results/loso_results.json
  data/results/interpolation_global_gmm/results/majority_vote_results.json

Writes:
  data/results/interpolation_global_gmm/results/loso_global_gmm_summary.xlsx

Run: python scripts/export_loso_global_gmm_excel.py
"""
import json
import re
import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

# Documented series time windows (CLAUDE.md) — used ONLY to disambiguate the
# raw filename date order, which is inconsistent across series: older series
# (10.01, 7.02, ...) embed dates as YYYY_DD_MM in places, newer ha-/hb-
# series (and some files within older series) use YYYY_MM_DD. Padding the
# window by a few days absorbs minor doc imprecision without creating
# ambiguity between the two orderings.
SERIES_WINDOWS = {
    '10.01.2026': ('2025-12-20', '2026-01-10'), '24.01.2026': ('2026-01-13', '2026-01-24'),
    '7.02.2026': ('2026-01-26', '2026-02-07'), '21.02.2026': ('2026-02-12', '2026-02-21'),
    '7.03.2026': ('2026-02-25', '2026-03-07'), '15.03.2026': ('2026-03-09', '2026-03-14'),
    '5.04.2026': ('2026-03-26', '2026-04-05'), '11.04.2026': ('2026-04-07', '2026-04-11'),
    '18.04.2026': ('2026-04-13', '2026-04-18'), 'ha-25.04.2026': ('2026-04-20', '2026-04-25'),
    'ha-2.05.2026': ('2026-04-27', '2026-05-02'), 'ha-17.05.2026': ('2026-05-12', '2026-05-17'),
    'ha-6.06.2026': ('2026-06-03', '2026-06-06'), 'ha-2.06.2026': ('2026-06-01', '2026-06-02'),
    'ha-20.06.2026': ('2026-06-17', '2026-06-20'),
}
WINDOW_PAD = datetime.timedelta(days=4)
MONTH_NAMES = {1: 'January', 2: 'February', 3: 'March', 4: 'April', 5: 'May',
               6: 'June', 7: 'July', 8: 'August', 9: 'September',
               10: 'October', 11: 'November', 12: 'December'}


def _try_date(y, mo, da):
    try:
        return datetime.date(y, mo, da)
    except ValueError:
        return None


def parse_file_date(series: str, fname: str):
    """
    Resolve a file's date, trying both YYYY_MM_DD and YYYY_DD_MM against the
    series' documented time window (padded) and picking whichever is valid.
    Returns None if neither order falls inside the window (shouldn't happen
    for the current 15 AUC-computable series — verified 2026-07-12).
    """
    m = re.search(r'(\d{4})_(\d{2})_(\d{2})', fname)
    if not m or series not in SERIES_WINDOWS:
        return None
    y, a, b = int(m.group(1)), int(m.group(2)), int(m.group(3))
    lo, hi = SERIES_WINDOWS[series]
    lo_d = datetime.date.fromisoformat(lo) - WINDOW_PAD
    hi_d = datetime.date.fromisoformat(hi) + WINDOW_PAD
    d_mmdd = _try_date(y, a, b)
    d_ddmm = _try_date(y, b, a)
    ok_mmdd = d_mmdd is not None and lo_d <= d_mmdd <= hi_d
    ok_ddmm = d_ddmm is not None and lo_d <= d_ddmm <= hi_d
    if ok_mmdd:
        return d_mmdd
    if ok_ddmm:
        return d_ddmm
    return None

ROOT = Path(__file__).resolve().parent.parent
RESULTS_DIR = ROOT / 'data' / 'results' / 'interpolation_global_gmm' / 'results'
OUT_PATH = RESULTS_DIR / 'loso_global_gmm_summary.xlsx'

with open(RESULTS_DIR / 'loso_results.json') as f:
    loso = json.load(f)
with open(RESULTS_DIR / 'majority_vote_results.json') as f:
    mv = json.load(f)

summary = loso.pop('_summary', {})

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(color='FFFFFF', bold=True)
BOLD = Font(bold=True)
THIN = Side(style='thin', color='D1D5DB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN = PatternFill('solid', fgColor='C6EFCE')
RED = PatternFill('solid', fgColor='FFC7CE')
GRAY = PatternFill('solid', fgColor='F3F4F6')
CENTER = Alignment(horizontal='center', vertical='center')


def style_header(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ═══════════════════════════════════════════════════════════════════════
# Sheet 1 — Overview
# ═══════════════════════════════════════════════════════════════════════
ov = wb.active
ov.title = 'Overview'
ov['A1'] = 'LOSO — Global GMM (corrected labels, DeepSAD Transformer)'
ov['A1'].font = Font(bold=True, size=14)
ov.merge_cells('A1:B1')

rows = [
    ('Mean ROC-AUC', round(summary.get('mean_roc_auc', float('nan')), 4)),
    ('Mean F1', round(summary.get('mean_f1', float('nan')), 4)),
    ('Series with computed AUC', summary.get('n_series')),
    ('Series skipped (no/NoFlow-only data)', sum(1 for v in loso.values() if not v)),
]
r = 3
for label, val in rows:
    ov.cell(row=r, column=1, value=label).font = BOLD
    ov.cell(row=r, column=2, value=val)
    r += 1

r += 1
ov.cell(row=r, column=1, value='Majority-vote file accuracy (all series)').font = BOLD
total_files = sum(len(files) for files in mv.values())
total_correct = sum(
    1 for files in mv.values() for v in files.values()
    if v['true_label'] == v['majority_pred']
)
ov.cell(row=r, column=2, value=f'{total_correct}/{total_files} ({100*total_correct/total_files:.1f}%)')
r += 2

ov.cell(row=r, column=1, value='Pooled confusion matrix (patch-level, across all AUC-computable series)').font = BOLD
r += 1
tp = sum(v['confusion_matrix']['tp'] for v in loso.values() if v)
tn = sum(v['confusion_matrix']['tn'] for v in loso.values() if v)
fp = sum(v['confusion_matrix']['fp'] for v in loso.values() if v)
fn = sum(v['confusion_matrix']['fn'] for v in loso.values() if v)
for label, val in [('TP (leakage correctly flagged)', tp), ('TN (NoFlow correctly cleared)', tn),
                    ('FP (NoFlow falsely flagged)', fp), ('FN (leakage missed)', fn)]:
    ov.cell(row=r, column=1, value=label)
    ov.cell(row=r, column=2, value=val)
    r += 1
ov.cell(row=r, column=1, value='Pooled accuracy').font = BOLD
ov.cell(row=r, column=2, value=round((tp + tn) / (tp + tn + fp + fn), 4))
r += 1
ov.cell(row=r, column=1, value='Pooled recall on leakage class').font = BOLD
ov.cell(row=r, column=2, value=round(tp / (tp + fn), 4))

autosize(ov, [42, 20])

# ═══════════════════════════════════════════════════════════════════════
# Sheet 2 — Per-series LOSO metrics
# ═══════════════════════════════════════════════════════════════════════
ws = wb.create_sheet('LOSO per-series')
cols = [
    'Series', 'ROC-AUC', 'F1', 'Threshold', 'N NoFlow', 'N Flow',
    'Accuracy', 'N Correct', 'N Wrong', 'TP', 'TN', 'FP', 'FN', 'Status',
]
for c, name in enumerate(cols, start=1):
    ws.cell(row=1, column=c, value=name)
style_header(ws, 1, len(cols))
ws.freeze_panes = 'A2'

row_i = 2
for series, v in loso.items():
    if not v:
        ws.cell(row=row_i, column=1, value=series)
        ws.cell(row=row_i, column=14, value='skipped (no data / NoFlow-only)')
        for c in range(1, len(cols) + 1):
            ws.cell(row=row_i, column=c).fill = GRAY
            ws.cell(row=row_i, column=c).border = BORDER
        row_i += 1
        continue
    cm = v['confusion_matrix']
    vals = [
        series, round(v['roc_auc'], 4), round(v['f1'], 4), round(v['threshold'], 2),
        v['n_noflow'], v['n_flow'], round(v['accuracy'], 4), v['n_correct'], v['n_wrong'],
        cm['tp'], cm['tn'], cm['fp'], cm['fn'],
        'inverted (AUC<0.5)' if v['roc_auc'] < 0.5 else ('weak' if v['roc_auc'] < 0.8 else 'ok'),
    ]
    for c, val in enumerate(vals, start=1):
        cell = ws.cell(row=row_i, column=c, value=val)
        cell.border = BORDER
    if v['roc_auc'] < 0.5:
        ws.cell(row=row_i, column=2).fill = RED
    row_i += 1

# Summary row
ws.cell(row=row_i, column=1, value='MEAN (AUC-computable series)').font = BOLD
ws.cell(row=row_i, column=2, value=round(summary.get('mean_roc_auc', float('nan')), 4)).font = BOLD
ws.cell(row=row_i, column=3, value=round(summary.get('mean_f1', float('nan')), 4)).font = BOLD
for c in range(1, len(cols) + 1):
    ws.cell(row=row_i, column=c).fill = HEADER_FILL if c <= 3 else GRAY
    if c <= 3:
        ws.cell(row=row_i, column=c).font = Font(color='FFFFFF', bold=True)

# Color scale on ROC-AUC column (red -> yellow -> green)
last_row = row_i - 1
ws.conditional_formatting.add(
    f'B2:B{last_row}',
    ColorScaleRule(
        start_type='min', start_color='FFC7CE',
        mid_type='num', mid_value=0.9, mid_color='FFEB9C',
        end_type='max', end_color='C6EFCE',
    ),
)

autosize(ws, [16, 10, 10, 12, 10, 9, 10, 10, 10, 8, 8, 8, 8, 22])

# ═══════════════════════════════════════════════════════════════════════
# Sheet 3 — Majority vote, per file, grouped by month
# ═══════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet('Majority Vote by Month')
cols2 = ['Date', 'Series', 'File', 'True Class', 'Predicted Class', 'Match',
         'Patches -> No Flow', 'Patches -> Flow', 'Total Patches']
N2 = len(cols2)


def class_label(raw: str) -> str:
    """LEAKAGE/NOFLOW -> Flow/No Flow, matching the user-facing wording."""
    return 'Flow' if raw == 'LEAKAGE' else 'No Flow'


# Build one row per file with a resolved date, then sort into month buckets.
rows_by_month = {}
unresolved = []
for series, files in mv.items():
    for fname, info in sorted(files.items()):
        d = parse_file_date(series, fname)
        if d is None:
            unresolved.append((series, fname, info))
            continue
        rows_by_month.setdefault((d.year, d.month), []).append((d, series, fname, info))

row_i = 1
month_row_ranges = []  # (start_row, end_row) of each month's data, for banding
for (year, month) in sorted(rows_by_month):
    file_rows = sorted(rows_by_month[(year, month)], key=lambda t: (t[1], t[0], t[2]))

    # Month divider (merged, bold, full width)
    ws2.cell(row=row_i, column=1, value=f'{MONTH_NAMES[month]} {year}  ({len(file_rows)} files)')
    ws2.cell(row=row_i, column=1).font = Font(bold=True, size=12, color='FFFFFF')
    ws2.cell(row=row_i, column=1).fill = HEADER_FILL
    ws2.merge_cells(start_row=row_i, start_column=1, end_row=row_i, end_column=N2)
    row_i += 1

    # Column header, repeated under each month divider so a printed/filtered
    # month section is self-contained.
    for c, name in enumerate(cols2, start=1):
        ws2.cell(row=row_i, column=c, value=name)
    style_header(ws2, row_i, N2)
    row_i += 1

    month_start = row_i
    n_flow = n_correct = 0
    for d, series, fname, info in file_rows:
        match = info['true_label'] == info['majority_pred']
        total = info['n_noflow'] + info['n_leakage']
        if info['true_label'] == 'LEAKAGE':
            n_flow += 1
        if match:
            n_correct += 1
        vals = [
            d.isoformat(), series, fname,
            class_label(info['true_label']), class_label(info['majority_pred']),
            'YES' if match else 'NO', info['n_noflow'], info['n_leakage'], total,
        ]
        for c, val in enumerate(vals, start=1):
            cell = ws2.cell(row=row_i, column=c, value=val)
            cell.border = BORDER
        fill = GREEN if match else RED
        for c in range(1, N2 + 1):
            ws2.cell(row=row_i, column=c).fill = fill
        row_i += 1
    month_row_ranges.append((month_start, row_i - 1))

    # Per-month subtotal row
    n_total = len(file_rows)
    ws2.cell(row=row_i, column=1, value='Subtotal').font = BOLD
    ws2.cell(row=row_i, column=4, value=f'{n_flow} Flow / {n_total - n_flow} No Flow').font = BOLD
    ws2.cell(row=row_i, column=6, value=f'{n_correct}/{n_total} correct').font = BOLD
    for c in range(1, N2 + 1):
        ws2.cell(row=row_i, column=c).fill = GRAY
        ws2.cell(row=row_i, column=c).border = BORDER
    row_i += 2  # blank row between months

if unresolved:
    ws2.cell(row=row_i, column=1, value=f'Unresolved date ({len(unresolved)} files) — see console output').font = BOLD
    row_i += 1
    for series, fname, info in unresolved:
        ws2.cell(row=row_i, column=2, value=series)
        ws2.cell(row=row_i, column=3, value=fname)
        row_i += 1
    print('[warn] could not resolve a date for:', unresolved)

ws2.freeze_panes = 'A1'
autosize(ws2, [12, 16, 34, 12, 14, 8, 16, 14, 14])

# ═══════════════════════════════════════════════════════════════════════
# Sheet 4 — Excluded / no-data series (so nothing is silently missing)
# ═══════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet('Excluded Series')
cols3 = ['Series', 'Reason']
for c, name in enumerate(cols3, start=1):
    ws3.cell(row=1, column=c, value=name)
style_header(ws3, 1, len(cols3))

excluded_reasons = {
    'ha-9.05.2026': (
        'Label status disputed — config has an alternating leakage pattern from a commit '
        'that flagged the raw data "possibly corrupted" and marked it processed:true to '
        'skip it. An earlier commit had relabelled all files no_leakage. No processed/'
        'prepared data exists on disk as of this run — contributes to no LOSO fold '
        'until a supervisor confirms the correct labels.'
    ),
    'ha-24.05.2026': (
        'Confirmed NoFlow-only series (0 leakage + 6 no_leakage) — AUC is mathematically '
        'undefined with only one class present, correctly excluded from the mean.'
    ),
}
r = 2
for series, reason in excluded_reasons.items():
    ws3.cell(row=r, column=1, value=series).font = BOLD
    ws3.cell(row=r, column=2, value=reason)
    ws3.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')
    ws3.row_dimensions[r].height = 60
    r += 1
autosize(ws3, [16, 100])

wb.save(OUT_PATH)
print(f'Saved -> {OUT_PATH}')
