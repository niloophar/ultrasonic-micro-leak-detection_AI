#!/usr/bin/env python
"""
Final file-level CNN vs. Transformer (DeepSAD) comparison table, requested
by Patrick. Both models are evaluated on the identical eval33 protocol:
train on 9 series (10.01, 24.01, 7.02, 7.03, 15.03, 5.04, 11.04, 18.04,
ha-25.04), score with a single global GMM (no per-site recalibration) on
3 held-out series (21.02, ha-2.05, ha-17.05). Neither model's checkpoint
is retrained or modified here — this script only reads the two saved
majority_vote_results_global.json files and reshapes them into thesis-ready
tables and a figure.

Reads:
  data/contrastive_g5_supcon_v2_ae_eval33/results/majority_vote_results_global.json
      (CNN, G5 SupCon AE dual mean-only pw64 noflow_filter, same 9/3 split)
  data/contrastive_g5_deepsad_transformer_ae_eval33/results/results/majority_vote_results_global.json
      (Transformer, DeepSAD, eval33 checkpoint, re-scored under gmm_eval_mode=global)

Writes:
  data/results/cnn_transformer_file_level_long_latex.txt       (main LaTeX table, long/narrow)
  data/results/cnn_transformer_patch_counts_appendix_latex.txt (appendix LaTeX table, wide)
  data/results/cnn_transformer_file_level_comparison.xlsx      (Summary / Long / Appendix sheets)
  figures/cnn_transformer_patch_bar_comparison.pdf
  figures/cnn_transformer_patch_bar_comparison.png

Run: python scripts/export_cnn_transformer_file_level_comparison.py
"""
import json
import re
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.transforms as mtransforms
from matplotlib.patches import Patch

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CNN_PATH = ROOT / 'data' / 'contrastive_g5_supcon_v2_ae_eval33' / 'results' / 'majority_vote_results_global.json'
TRF_PATH = ROOT / 'data' / 'contrastive_g5_deepsad_transformer_ae_eval33' / 'results' / 'results' / 'majority_vote_results_global.json'
RESULTS_OUT = ROOT / 'data' / 'results'
FIG_DIR = ROOT / 'figures'
RESULTS_OUT.mkdir(parents=True, exist_ok=True)
FIG_DIR.mkdir(parents=True, exist_ok=True)

SERIES_ORDER = ['21.02.2026', 'ha-2.05.2026', 'ha-17.05.2026']
SERIES_LABEL = {'21.02.2026': '21.02', 'ha-2.05.2026': 'ha-2.05', 'ha-17.05.2026': 'ha-17.05'}
LABEL_MAP = {'NOFLOW': 'NoFlow', 'LEAKAGE': 'Leakage'}
DATE_RE = re.compile(r'(\d{4})_(\d{2})_(\d{2})')

# ─── calm, colorblind-safe categorical pair (validated: node validate_palette.js
# "#2a78d6,#eb6834" --mode light -> all checks pass) ────────────────────────
NF_COLOR = '#2a78d6'    # blue  — NoFlow
LEAK_COLOR = '#eb6834'  # orange — Leakage
C_DARK = '#1a1a1a'
C_GRAY = '#777777'
C_LGRAY = '#cccccc'


def file_date(fname: str) -> str:
    m = DATE_RE.search(fname)
    return f'{m.group(1)}-{m.group(2)}-{m.group(3)}' if m else fname


with open(CNN_PATH) as f:
    cnn = json.load(f)
with open(TRF_PATH) as f:
    trf = json.load(f)

# ─── build one unified, ordered record list ────────────────────────────────
groups = []
for series in SERIES_ORDER:
    for fname in sorted(cnn.get(series, {}).keys()):
        c = cnn[series][fname]
        t = trf[series][fname]
        assert c['true_label'] == t['true_label'], (
            f'GT mismatch for {series}/{fname}: CNN={c["true_label"]} vs Transformer={t["true_label"]}'
        )
        groups.append({
            'series': series, 'fname': fname, 'date': file_date(fname),
            'gt': LABEL_MAP[c['true_label']], 'cnn': c, 'trf': t,
        })

print(f'{len(groups)} files across {len(SERIES_ORDER)} series loaded.')

# Long-format rows (one row = one model applied to one file)
long_rows = []
for g in groups:
    for model, info in (('CNN', g['cnn']), ('Transformer', g['trf'])):
        n = info['n_noflow'] + info['n_leakage']
        mv = LABEL_MAP[info['majority_pred']]
        long_rows.append({
            'series': SERIES_LABEL[g['series']], 'date': g['date'], 'model': model,
            'gt': g['gt'], 'n': n, 'nf': info['n_noflow'], 'leak': info['n_leakage'],
            'mv': mv, 'correct': mv == g['gt'],
        })

# Wide/appendix rows (one row = one file, both models' counts side by side)
wide_rows = []
for g in groups:
    n_cnn = g['cnn']['n_noflow'] + g['cnn']['n_leakage']
    n_trf = g['trf']['n_noflow'] + g['trf']['n_leakage']
    assert n_cnn == n_trf, f'Patch-count mismatch for {g["series"]}/{g["fname"]}: CNN N={n_cnn} vs Transformer N={n_trf}'
    mv_cnn = LABEL_MAP[g['cnn']['majority_pred']]
    mv_trf = LABEL_MAP[g['trf']['majority_pred']]
    wide_rows.append({
        'series': SERIES_LABEL[g['series']], 'date': g['date'], 'gt': g['gt'], 'n': n_cnn,
        'cnn_nf': g['cnn']['n_noflow'], 'cnn_leak': g['cnn']['n_leakage'], 'cnn_mv': mv_cnn,
        'cnn_correct': mv_cnn == g['gt'],
        'trf_nf': g['trf']['n_noflow'], 'trf_leak': g['trf']['n_leakage'], 'trf_mv': mv_trf,
        'trf_correct': mv_trf == g['gt'],
    })

# ═══════════════════════════════════════════════════════════════════════
# LaTeX — main long/narrow table
# ═══════════════════════════════════════════════════════════════════════
lines = []
lines.append(r'\begin{longtable}{lllrrrl}')
lines.append(
    r'\caption{File-level majority-vote comparison between the CNN and '
    r'Transformer (DeepSAD) encoders under the deployment-realistic global-GMM '
    r'protocol (train: 9 series excl.\ 21.02; test: 21.02, ha-2.05, ha-17.05; '
    r'no per-site recalibration). $N$ is the total number of evaluated patches '
    r'in the file; Leak is the number of those patches predicted Leakage. '
    r'A dagger ($\dagger$) marks a majority-vote decision that disagrees with '
    r'the ground truth.}'
)
lines.append(r'\label{tab:cnn_transformer_file_level}\\')
lines.append(r'\toprule')
lines.append(r'Series & File/day & Model & GT & $N$ & Leak patches & MV Decision \\')
lines.append(r'\midrule')
lines.append(r'\endfirsthead')
lines.append(r'\toprule')
lines.append(r'Series & File/day & Model & GT & $N$ & Leak patches & MV Decision \\')
lines.append(r'\midrule')
lines.append(r'\endhead')
lines.append(r'\midrule')
lines.append(r'\multicolumn{7}{r}{\textit{continued on next page}} \\')
lines.append(r'\midrule')
lines.append(r'\endfoot')
lines.append(r'\bottomrule')
lines.append(r'\endlastfoot')

prev_series = None
for i, row in enumerate(long_rows):
    if prev_series is not None and row['series'] != prev_series:
        lines.append(r'\midrule')
    elif prev_series is not None and i % 2 == 0:
        lines.append(r'\addlinespace[2pt]')
    prev_series = row['series']
    mv_txt = row['mv'] + ('' if row['correct'] else r'$^\dagger$')
    lines.append(
        f"{row['series']} & {row['date']} & {row['model']} & {row['gt']} & "
        f"{row['n']} & {row['leak']} & {mv_txt} \\\\"
    )

lines.append(r'\end{longtable}')
lines.append('')
lines.append(r'$\dagger$ majority-vote decision differs from ground truth.')

main_latex_path = RESULTS_OUT / 'cnn_transformer_file_level_long_latex.txt'
main_latex_path.write_text('\n'.join(lines) + '\n')
print(f'Saved -> {main_latex_path}')

# ═══════════════════════════════════════════════════════════════════════
# LaTeX — appendix wide table (patch-count detail, one row per file)
# ═══════════════════════════════════════════════════════════════════════
alines = []
alines.append(r'\begin{longtable}{llcr rrl rrl}')
alines.append(
    r'\caption{Per-file patch counts for both models under the same global-GMM '
    r'protocol as \cref{tab:cnn_transformer_file_level}, shown side by side for '
    r'direct comparison. NF/Leak are patch counts predicted NoFlow/Leakage; MV '
    r'is the resulting majority-vote file decision. A dagger ($\dagger$) marks a '
    r'majority-vote decision that disagrees with the ground truth.}'
)
alines.append(r'\label{tab:cnn_transformer_appendix}\\')
alines.append(r'\toprule')
header = (
    r'Series & File/day & GT & $N$ & '
    r'\thead{CNN\\NF} & \thead{CNN\\Leak} & \thead{CNN\\MV} & '
    r'\thead{Transformer\\NF} & \thead{Transformer\\Leak} & \thead{Transformer\\MV} \\'
)
alines.append(header)
alines.append(r'\midrule')
alines.append(r'\endfirsthead')
alines.append(r'\toprule')
alines.append(header)
alines.append(r'\midrule')
alines.append(r'\endhead')
alines.append(r'\midrule')
alines.append(r'\multicolumn{10}{r}{\textit{continued on next page}} \\')
alines.append(r'\midrule')
alines.append(r'\endfoot')
alines.append(r'\bottomrule')
alines.append(r'\endlastfoot')

prev_series = None
for row in wide_rows:
    if prev_series is not None and row['series'] != prev_series:
        alines.append(r'\midrule')
    prev_series = row['series']
    cnn_mv_txt = row['cnn_mv'] + ('' if row['cnn_correct'] else r'$^\dagger$')
    trf_mv_txt = row['trf_mv'] + ('' if row['trf_correct'] else r'$^\dagger$')
    alines.append(
        f"{row['series']} & {row['date']} & {row['gt']} & {row['n']} & "
        f"{row['cnn_nf']} & {row['cnn_leak']} & {cnn_mv_txt} & "
        f"{row['trf_nf']} & {row['trf_leak']} & {trf_mv_txt} \\\\"
    )

alines.append(r'\end{longtable}')
alines.append('')
alines.append(r'$\dagger$ majority-vote decision differs from ground truth.')

appendix_latex_path = RESULTS_OUT / 'cnn_transformer_patch_counts_appendix_latex.txt'
appendix_latex_path.write_text('\n'.join(alines) + '\n')
print(f'Saved -> {appendix_latex_path}')

# ═══════════════════════════════════════════════════════════════════════
# Excel workbook — Summary / Long format / Appendix wide format
# ═══════════════════════════════════════════════════════════════════════
HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(color='FFFFFF', bold=True)
BOLD = Font(bold=True)
THIN = Side(style='thin', color='D1D5DB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN = PatternFill('solid', fgColor='C6EFCE')
RED = PatternFill('solid', fgColor='FFC7CE')
GRAY = PatternFill('solid', fgColor='F3F4F6')
CENTER = Alignment(horizontal='center', vertical='center')


def style_header(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# --- Sheet 1: Summary --------------------------------------------------
ov = wb.active
ov.title = 'Summary'
ov['A1'] = 'CNN vs. Transformer (DeepSAD) — File-Level Comparison (eval33 protocol, global GMM)'
ov['A1'].font = Font(bold=True, size=13)
ov.merge_cells('A1:D1')
ov['A2'] = 'Train: 10.01, 24.01, 7.02, 7.03, 15.03, 5.04, 11.04, 18.04, ha-25.04 (9 series)  |  Test: 21.02, ha-2.05, ha-17.05'
ov['A2'].font = Font(italic=True, size=9, color='555555')
ov.merge_cells('A2:D2')

r = 4
cols_s = ['Series', 'N Files', 'CNN correct', 'Transformer correct']
for c, name in enumerate(cols_s, start=1):
    ov.cell(row=r, column=c, value=name)
style_header(ov, r, len(cols_s))
r += 1
totals = {'n': 0, 'cnn': 0, 'trf': 0}
for series in SERIES_ORDER:
    s_rows = [w for w in wide_rows if w['series'] == SERIES_LABEL[series]]
    n_files = len(s_rows)
    cnn_correct = sum(1 for w in s_rows if w['cnn_correct'])
    trf_correct = sum(1 for w in s_rows if w['trf_correct'])
    totals['n'] += n_files
    totals['cnn'] += cnn_correct
    totals['trf'] += trf_correct
    vals = [SERIES_LABEL[series], n_files, f'{cnn_correct}/{n_files}', f'{trf_correct}/{n_files}']
    for c, v in enumerate(vals, start=1):
        cell = ov.cell(row=r, column=c, value=v)
        cell.border = BORDER
    r += 1
ov.cell(row=r, column=1, value='TOTAL').font = BOLD
ov.cell(row=r, column=2, value=totals['n']).font = BOLD
ov.cell(row=r, column=3, value=f"{totals['cnn']}/{totals['n']} ({100*totals['cnn']/totals['n']:.1f}%)").font = BOLD
ov.cell(row=r, column=4, value=f"{totals['trf']}/{totals['n']} ({100*totals['trf']/totals['n']:.1f}%)").font = BOLD
for c in range(1, 5):
    ov.cell(row=r, column=c).fill = GRAY
    ov.cell(row=r, column=c).border = BORDER
autosize(ov, [14, 10, 16, 20])

# --- Sheet 2: Long format ----------------------------------------------
ws = wb.create_sheet('File-Level Long')
cols_l = ['Series', 'File/day', 'Model', 'GT', 'N', 'NoFlow patches', 'Leak patches', 'MV Decision', 'Correct']
for c, name in enumerate(cols_l, start=1):
    ws.cell(row=1, column=c, value=name)
style_header(ws, 1, len(cols_l))
ws.freeze_panes = 'A2'
for r_i, row in enumerate(long_rows, start=2):
    vals = [row['series'], row['date'], row['model'], row['gt'], row['n'], row['nf'], row['leak'],
            row['mv'], 'Yes' if row['correct'] else 'No']
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r_i, column=c, value=v)
        cell.border = BORDER
    fill = GREEN if row['correct'] else RED
    for c in range(1, len(cols_l) + 1):
        ws.cell(row=r_i, column=c).fill = fill
autosize(ws, [12, 12, 13, 10, 8, 14, 13, 13, 9])

# --- Sheet 3: Appendix wide format --------------------------------------
ws2 = wb.create_sheet('Patch Counts Appendix')
cols_w = ['Series', 'File/day', 'GT', 'N', 'CNN NF', 'CNN Leak', 'CNN MV', 'CNN Correct',
          'Transformer NF', 'Transformer Leak', 'Transformer MV', 'Transformer Correct']
for c, name in enumerate(cols_w, start=1):
    ws2.cell(row=1, column=c, value=name)
style_header(ws2, 1, len(cols_w))
ws2.freeze_panes = 'A2'
for r_i, row in enumerate(wide_rows, start=2):
    vals = [row['series'], row['date'], row['gt'], row['n'],
            row['cnn_nf'], row['cnn_leak'], row['cnn_mv'], 'Yes' if row['cnn_correct'] else 'No',
            row['trf_nf'], row['trf_leak'], row['trf_mv'], 'Yes' if row['trf_correct'] else 'No']
    for c, v in enumerate(vals, start=1):
        cell = ws2.cell(row=r_i, column=c, value=v)
        cell.border = BORDER
    ws2.cell(row=r_i, column=7).fill = GREEN if row['cnn_correct'] else RED
    ws2.cell(row=r_i, column=11).fill = GREEN if row['trf_correct'] else RED
autosize(ws2, [12, 12, 10, 8, 9, 10, 10, 11, 14, 15, 15, 17])

xlsx_path = RESULTS_OUT / 'cnn_transformer_file_level_comparison.xlsx'
wb.save(xlsx_path)
print(f'Saved -> {xlsx_path}')

# ═══════════════════════════════════════════════════════════════════════
# Figure — compact 3-panel layout, one panel per series: (a) 21.02,
# (b) ha-2.05, (c) ha-17.05. Each panel: horizontal stacked bars, one
# CNN + one Transformer bar per file/day. Independent x-scale per panel
# (patch counts differ by an order of magnitude across series — a shared
# scale would flatten the smaller series to illegibility).
# ═══════════════════════════════════════════════════════════════════════
Y_STEP = 1.0
BAR_H = 0.34
OFFSET = 0.20
PANEL_LETTERS = {'21.02.2026': '(a)', 'ha-2.05.2026': '(b)', 'ha-17.05.2026': '(c)'}

panels = {s: [g for g in groups if g['series'] == s] for s in SERIES_ORDER}
max_n = max(len(v) for v in panels.values())

fig_height = max_n * 0.62 + 1.5
fig, axes = plt.subplots(1, 3, figsize=(13.5, fig_height))

for ax, series in zip(axes, SERIES_ORDER):
    sub = panels[series]
    yaxis_frac = ax.get_yaxis_transform()  # x in axes-fraction, y in data coords

    for i, g in enumerate(sub):
        y_cnn = i * Y_STEP + OFFSET
        y_trf = i * Y_STEP - OFFSET
        for y_pos, info in ((y_cnn, g['cnn']), (y_trf, g['trf'])):
            nf, leak = info['n_noflow'], info['n_leakage']
            ax.barh(y_pos, nf, left=0, height=BAR_H, color=NF_COLOR,
                    edgecolor='white', linewidth=0.7, zorder=3)
            ax.barh(y_pos, leak, left=nf, height=BAR_H, color=LEAK_COLOR,
                    edgecolor='white', linewidth=0.7, zorder=3)
        ax.text(-0.02, y_cnn, 'CNN', transform=yaxis_frac, ha='right', va='center',
                fontsize=7.0, color=C_GRAY)
        ax.text(-0.02, y_trf, 'Transf.', transform=yaxis_frac, ha='right', va='center',
                fontsize=7.0, color=C_GRAY)
        short_date = g['date'][5:]  # 'YYYY-MM-DD' -> 'MM-DD' (year is 2026 throughout)
        ax.text(-0.19, i * Y_STEP, f"{short_date}\n({g['gt']})", transform=yaxis_frac,
                ha='right', va='center', fontsize=8.0, color=C_DARK, linespacing=1.25)

    ax.set_title(f'{PANEL_LETTERS[series]} {SERIES_LABEL[series]}', loc='left',
                 fontsize=11.5, fontweight='bold', color=C_DARK, pad=8)
    ax.set_ylim(-0.5, max_n - 1 + 0.5)
    ax.invert_yaxis()
    ax.set_yticks([])
    ax.tick_params(axis='x', colors=C_DARK, labelsize=8.5)
    ax.grid(axis='x', color=C_LGRAY, lw=0.5, alpha=0.5, zorder=0)
    for spine in ('top', 'right', 'left'):
        ax.spines[spine].set_visible(False)
    ax.spines['bottom'].set_color(C_LGRAY)

legend_handles = [
    Patch(facecolor=NF_COLOR, edgecolor='white', label='Predicted NoFlow'),
    Patch(facecolor=LEAK_COLOR, edgecolor='white', label='Predicted Leakage'),
]
fig.legend(handles=legend_handles, loc='upper center', bbox_to_anchor=(0.5, 1.0),
           ncol=2, frameon=False, fontsize=9.5)

fig.suptitle('CNN vs. Transformer (DeepSAD) — File-Level Patch Classification',
             fontsize=12.5, fontweight='bold', color=C_DARK, y=1.045)
fig.text(
    0.5, -0.05,
    'Global GMM, no per-site recalibration. Train: 10.01, 24.01, 7.02, 7.03, 15.03, 5.04, 11.04, 18.04, '
    'ha-25.04 (9 series). Ground truth shown in parentheses next to each date. Patch axis independent per panel.',
    ha='center', va='top', fontsize=7.5, color=C_GRAY, linespacing=1.35, transform=fig.transFigure,
)
fig.text(0.5, -0.005, 'Number of evaluated patches', ha='center', va='top',
         fontsize=9.5, color=C_DARK, transform=fig.transFigure)

fig.subplots_adjust(left=0.09, right=0.99, bottom=0.14, top=0.86, wspace=0.55)
for ext in ('png', 'pdf'):
    kw = dict(format=ext, bbox_inches='tight')
    if ext == 'png':
        kw['dpi'] = 300
    out_path = FIG_DIR / f'cnn_transformer_patch_bar_comparison.{ext}'
    fig.savefig(out_path, **kw)
    print(f'Saved -> {out_path}')
