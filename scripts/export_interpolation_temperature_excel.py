#!/usr/bin/env python
"""
Export the interpolation temperature comparison (no-temp / pipe / pipe+ambient)
to a single formatted Excel workbook for the thesis writeup.

Reads:
  data/results/interpolation_temperature/results/test_metrics_global.json
  data/results/interpolation_temperature/results/test_metrics_global_conditional.json
  data/results/interpolation_temperature/results/test_metrics_global_conditional_temp_pipe_temp_amb.json

Writes:
  data/results/interpolation_temperature/results/interpolation_temperature_summary.xlsx

Run: python scripts/export_interpolation_temperature_excel.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
RESULTS_DIR = ROOT / 'data' / 'results' / 'interpolation_temperature' / 'results'
OUT_PATH = RESULTS_DIR / 'interpolation_temperature_summary.xlsx'

TEST_SERIES = 'ha-2.06.2026'
SETTINGS = [
    ('No Temperature', 'test_metrics_global.json'),
    ('Pipe Temperature', 'test_metrics_global_conditional.json'),
    ('Pipe + Ambient', 'test_metrics_global_conditional_temp_pipe_temp_amb.json'),
]

loaded = {}
for label, fname in SETTINGS:
    with open(RESULTS_DIR / fname) as f:
        data = json.load(f)
    loaded[label] = {'series': data[TEST_SERIES], 'overall': data['overall']}

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(color='FFFFFF', bold=True)
BOLD = Font(bold=True)
THIN = Side(style='thin', color='D1D5DB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GRAY = PatternFill('solid', fgColor='F3F4F6')
YELLOW = PatternFill('solid', fgColor='FFF9C4')
CENTER = Alignment(horizontal='center', vertical='center')


def style_header(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ═══════════════════════════════════════════════════════════════════════
# Sheet 1 — Comparison (main deliverable: metric rows x setting columns)
# ═══════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = 'Comparison'
ws['A1'] = f'Interpolation Temperature Comparison — test series {TEST_SERIES}'
ws['A1'].font = Font(bold=True, size=14)
ws.merge_cells('A1:E1')
ws['A2'] = '15-series train (incl. ha-6.06.2026), same split as the ha_interp experiment. DeepSAD Transformer, Global GMM.'
ws['A2'].font = Font(italic=True, color='6B7280')
ws.merge_cells('A2:E2')

cols = ['Metric'] + [label for label, _ in SETTINGS] + ['Δ (Pipe+Ambient vs No Temp)']
header_row = 4
for c, name in enumerate(cols, start=1):
    ws.cell(row=header_row, column=c, value=name)
style_header(ws, header_row, len(cols))

def g(label, *keys):
    d = loaded[label]['series']
    for k in keys:
        d = d[k]
    return d

def go(label, key):
    return loaded[label]['overall'].get(key)

rows_spec = [
    ('ROC-AUC', lambda l: round(g(l, 'roc_auc'), 4), True),
    ('F1', lambda l: round(g(l, 'f1'), 4), True),
    ('Accuracy', lambda l: round(g(l, 'accuracy'), 4), True),
    ('Threshold', lambda l: round(g(l, 'threshold'), 2), False),
    ('N NoFlow (test)', lambda l: g(l, 'n_noflow'), False),
    ('N Flow (test)', lambda l: g(l, 'n_flow'), False),
    ('N Correct', lambda l: g(l, 'n_correct'), False),
    ('N Wrong', lambda l: g(l, 'n_wrong'), False),
    ('TP', lambda l: g(l, 'confusion_matrix', 'tp'), False),
    ('TN', lambda l: g(l, 'confusion_matrix', 'tn'), False),
    ('FP', lambda l: g(l, 'confusion_matrix', 'fp'), False),
    ('FN', lambda l: g(l, 'confusion_matrix', 'fn'), False),
    ('temp_pipe (°C, series mean)', lambda l: round(g(l, 'temp_pipe'), 2) if 'temp_pipe' in loaded[l]['series'] else '—', False),
    ('temp_amb (°C, series mean)', lambda l: round(g(l, 'temp_amb'), 2) if 'temp_amb' in loaded[l]['series'] else '—', False),
    ('GMM fit N (train NoFlow)', lambda l: go(l, 'gmm_fit_n_noflow'), False),
    ('slope_norm (regression magnitude)', lambda l: round(go(l, 'slope_norm'), 3) if go(l, 'slope_norm') is not None else '—', False),
]

r = header_row + 1
auc_row = f1_row = acc_row = None
for name, fn, highlight in rows_spec:
    ws.cell(row=r, column=1, value=name).font = BOLD if highlight else Font()
    vals = []
    for label, _ in SETTINGS:
        v = fn(label)
        vals.append(v)
        c = 2 + [l for l, _ in SETTINGS].index(label)
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = BORDER
        if highlight:
            cell.font = BOLD
            cell.fill = YELLOW
    # delta column: pipe+ambient vs no-temp, numeric rows only
    try:
        delta = round(vals[2] - vals[0], 4)
        ws.cell(row=r, column=len(cols), value=delta)
    except TypeError:
        ws.cell(row=r, column=len(cols), value='—')
    for c in range(1, len(cols) + 1):
        ws.cell(row=r, column=c).border = BORDER
    if name == 'ROC-AUC':
        auc_row = r
    r += 1

r += 1
ws.cell(row=r, column=1, value='Conclusion').font = Font(bold=True, size=12)
r += 1
conclusion = (
    "NULL RESULT: temperature conditioning does not improve global-GMM scoring on this "
    "series. AUC is flat to within 0.0013 across all three settings; Pipe-only and "
    "Pipe+Ambient produce an IDENTICAL confusion matrix (TP=7, TN=209, FP=3, FN=18) — "
    "adding temp_amb on top of temp_pipe changed nothing at all. This confirms the "
    "mechanism already established in Eval C (CLAUDE.md): AUC is a within-series ranking "
    "metric, and linear mean-conditioning shifts NoFlow and Flow scores by the same "
    "amount, leaving the ranking (and therefore AUC) unchanged — regardless of encoder "
    "architecture (CNN -> Transformer) or number of temperature covariates (1 -> 2)."
)
ws.cell(row=r, column=1, value=conclusion)
ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
ws.row_dimensions[r].height = 90

r += 2
ws.cell(row=r, column=1, value='Caveats found during analysis').font = Font(bold=True, size=12)
r += 1
caveats = [
    "GMM fit population differs between rows: 'No Temperature' was produced via --run "
    "train (fits on train_ids only, 7929 NoFlow patches); 'Pipe' and 'Pipe+Ambient' were "
    "produced via --run eval reusing the same checkpoint (fits on train_ids+val_ids, "
    "8949 NoFlow patches). This is an existing inconsistency between train_contrastive() "
    "and eval_contrastive(), not something this experiment introduced. It doesn't affect "
    "the Pipe vs Pipe+Ambient comparison (both went through eval_contrastive identically), "
    "but the No-Temperature row isn't fit on the exact same pool as the other two.",
    "Run-to-run variance: this same held-out series (ha-2.06.2026) scored AUC 0.072 "
    "(inverted) in the systematic 17-fold LOSO Global-GMM run three days earlier, vs. "
    "0.958 here, on a near-identical training pool. A single global-GMM AUC on an "
    "out-of-envelope series should be treated as noisy, not as a stable property of "
    "'the encoder', unless averaged over multiple training seeds.",
]
for cav in caveats:
    ws.cell(row=r, column=1, value='• ' + cav)
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
    ws.row_dimensions[r].height = 75
    r += 1

ws.freeze_panes = 'A5'
autosize(ws, [32, 18, 18, 18, 22])

wb.save(OUT_PATH)
print(f'Saved -> {OUT_PATH}')
